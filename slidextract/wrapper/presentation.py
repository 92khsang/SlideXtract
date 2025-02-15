from __future__ import annotations

import weakref
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from pptx import Presentation

from slidextract.wrapper.slide import SlideWrapper

if TYPE_CHECKING:
    from slidextract.wrapper.slide import SlideFilter


@dataclass(frozen=True)
class PresentationWrapper:
    """
    A wrapper for a PPTX presentation.

    Attributes:
        pptx_path (path): The path to the PPTX file.
        slide_filter (SlideFilter): Filter conditions for extracting slides
        presentation (Presentation): The original PPTX presentation
        slides (list[SlideWrapper]): List of extracted slides
    """

    __slots__ = ("__weakref__", "pptx_path", "slide_filter", "presentation", "slides")

    pptx_path: Path
    slide_filter: SlideFilter

    presentation: Presentation
    slides: list[SlideWrapper]

    def __init__(self, pptx_path: str | Path, slide_filter: SlideFilter):
        pptx_path = Path(pptx_path)
        if not pptx_path.is_file() or pptx_path.suffix != ".pptx":
            raise ValueError(f"Invalid PowerPoint file: {pptx_path}")

        object.__setattr__(self, "pptx_path", pptx_path)
        object.__setattr__(self, "slide_filter", slide_filter)
        object.__setattr__(self, "presentation", Presentation(str(self.pptx_path)))
        object.__setattr__(
            self,
            "slides",
            [
                SlideWrapper(
                    slide=slide,
                    number=self.presentation.slides.index(slide) + 1,
                    slide_filter=self.slide_filter,
                    _presentation_proxy=weakref.proxy(self),
                )
                for slide in self.presentation.slides
            ],
        )

    def __getattr__(self, item):
        return getattr(self.presentation, item)

    @property
    def total_slides(self) -> int:
        return len(self.slides)
