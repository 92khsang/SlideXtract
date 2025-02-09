from dataclasses import dataclass, field
from typing import TypeAlias, TYPE_CHECKING

from slidextract.wrapper.models import SlideSize, BBox

if TYPE_CHECKING:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.shapetree import _BaseGroupShapes, GroupShapes
    from pptx.shapes.autoshape import Shape as AutoShape
    from pptx.shapes.base import BaseShape
    from pptx.shapes.graphfrm import GraphicFrame
    from pptx.shapes.group import GroupShape
    from pptx.shapes.picture import Picture
    from pptx.slide import Slide

    Shape: TypeAlias = BaseShape | GroupShape | GraphicFrame | Picture | AutoShape


@dataclass(frozen=True, slots=True)
class ShapeFilter:
    """
    Filter conditions for extracting shapes.

    Attributes:
        min_width: minimum width of the shape
        min_height: minimum height of a shape
        exclude_types: tuple of shape types to exclude
        include_empty_textbox: whether to include empty text boxes
    """

    min_width: int = 0
    min_height: int = 0
    exclude_types: tuple[MSO_SHAPE_TYPE, ...] = ()
    include_empty_textbox: bool = True


@dataclass(frozen=True, slots=True)
class ShapeWrapper:
    """
    Wrapper for python-pptx shape

    Attributes:
        shape: Original shape
        slide_size: Size of the slide
        shape_filter: Filter conditions for extracting shape
        bbox: Bounding box of the shape
        valid: Whether the shape is valid
    """

    shape: Shape
    slide_size: SlideSize
    shape_filter: ShapeFilter
    bbox: BBox = field(init=False)
    valid: bool = field(init=False)

    _shapes: _BaseGroupShapes = field(init=False)
    _parent: Shape | Slide = field(init=False)

    def __post_init__(self):
        object.__setattr__(self, "_shapes", getattr(self.shape, "_parent", None))
        object.__setattr__(self, "_parent", getattr(self._shapes, "_parent", None))

        object.__setattr__(self, "bbox", self._calculate_bbox())
        object.__setattr__(self, "valid", self._validate())

    def _validate(self):
        return (
            self.shape_filter.min_width <= self.bbox.width
            and self.shape_filter.min_height <= self.bbox.height
            and self.shape.shape_type not in self.shape_filter.exclude_types
            and (
                self.shape_filter.include_empty_textbox
                or not self.type == MSO_SHAPE_TYPE.TEXT_BOX
                or self.has_text
            )
        )

    def _calculate_bbox(self) -> BBox:
        offset_x, offset_y, scale_x, scale_y = self._calculate_group_factors()

        left = max(0, (self.shape.left + offset_x))
        top = max(0, (self.shape.top + offset_y))

        width = int(float(self.shape.width) * scale_x)
        height = int(float(self.shape.height) * scale_y)

        right = min(self.slide_size.width, left + width)
        bottom = min(self.slide_size.height, top + height)

        return BBox(left, top, right, bottom)

    def _calculate_group_factors(self) -> tuple[int, int, float, float]:
        offset_x, offset_y = 0, 0
        scale_x, scale_y = 1.0, 1.0

        if self._shapes:
            if isinstance(self._shapes, GroupShapes) and self._parent:
                group = self._parent
                group_element = getattr(group, "_element")

                group_offset_x = group.left
                group_offset_y = group.top

                ch_off = group_element.xpath("./p:grpSpPr/a:xfrm/a:chOff")[0]
                ch_off_x = int(ch_off.get("x"))
                ch_off_y = int(ch_off.get("y"))

                group_width = group.width
                group_height = group.height

                ch_ext = group_element.xpath("./p:grpSpPr/a:xfrm/a:chExt")[0]
                ch_ext_x = int(ch_ext.get("cx"))
                ch_ext_y = int(ch_ext.get("cy"))

                if ch_ext_x > 0 and ch_ext_y > 0:
                    scale_x = group_width / ch_ext_x
                    scale_y = group_height / ch_ext_y

                offset_x = group_offset_x - ch_off_x
                offset_y = group_offset_y - ch_off_y

        return offset_x, offset_y, scale_x, scale_y

    @property
    def type(self):
        return self.shape.shape_type

    @property
    def has_text(self):
        return self.shape.has_text_frame and self.shape.text.strip() != ""
